"""Генерация Excel-таблиц оргструктуры из раздела 3 контрольных документов.

ТЗ называет среди входов «организационные структуры» в Excel. У нас таких файлов нет,
поэтому строим их из того же источника, что и docx: раздел 3 положения. Это честный
тест парсера Excel: структурные выводы из xlsx обязаны совпасть со структурными
выводами из docx, и это проверяется в test_formats.py.

    uv run make_testdata.py
"""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from orgdiff.parse import departments, parse

HERE = Path(__file__).parent
TZ = HERE.parent / "TZ"
OUT = HERE / "testdata"

SUBORD = re.compile(r"^(?:Директору|Главному аудитору)\s*([А-ЯЁа-яё\s]*?)\s*подчиняются", re.I)


def build(docx: Path, xlsx: Path, doc: str):
    clauses = parse(str(docx), doc)
    depts = departments(clauses)
    by_id = {c.id: c for c in clauses}

    wb = Workbook()
    ws = wb.active
    ws.title = "Структура"
    ws.append(["Подразделение", "Аббревиатура", "Должность", "Подчиняется", "Источник"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    # подразделения из 3.4
    for abbr, full in depts.items():
        src = next((c.id for c in clauses if c.num == "3.4" and abbr in c.text), "3.4")
        ws.append([full, abbr, "", "БВА", src])

    # должности: подпункты 3.5-3.9, владелец из вводной фразы "Директору X подчиняются ..."
    for c in clauses:
        if c.section != "3" or not c.num or c.num == "3.4" or c.id == c.num:
            continue
        lead = by_id.get(c.num)
        if not lead or not lead.is_lead:
            continue
        m = SUBORD.match(lead.text)
        if not m:
            continue
        boss_raw = m.group(1).strip()
        boss = "Главный аудитор" if lead.text.startswith("Главному") else f"Директор {boss_raw}".strip()
        unit_abbr = next((a for a in depts if a in boss_raw), "")
        unit_full = depts.get(unit_abbr, "БВА")
        ws.append([unit_full, unit_abbr, c.text.rstrip("."), boss, c.id])

    for col, width in zip("ABCDE", (52, 14, 48, 34, 10), strict=True):
        ws.column_dimensions[col].width = width
    OUT.mkdir(exist_ok=True)
    wb.save(xlsx)
    print(f"{xlsx.name}: {ws.max_row - 1} строк")


if __name__ == "__main__":
    for n in (8, 9):
        build(TZ / f"Положение_о_внутреннем_аудите_редакция_{n}_обезличено.docx", OUT / f"структура_ред{n}.xlsx", f"ред.{n}")

"""Итоговое аналитическое заключение: из результата конвейера в markdown и docx.

Must have 5 из ТЗ: «сформировать итоговое аналитическое заключение в понятном для
пользователя виде». Структура заключения повторяет раздел «Вход -> выход» ТЗ:
перечень подразделений, таблица сопоставления, потери, дубли, конфликты интересов,
ссылки на пункты, краткие рекомендации. Рекомендации генерирует модель по строгой
схеме и только со ссылкой на конкретный вывод; без ключа ставится шаблонная
формулировка. Всё заключение носит рекомендательный характер, это написано в нём самом.
"""

from __future__ import annotations

import json
import os
from datetime import date

LOSS_KINDS = ("lost", "partial_loss", "removed_from_owner")
MOVE_KINDS = ("moved_to", "moved_from", "merged", "reassigned", "covered", "added_to_owner")
NEW_KINDS = ("new", "partial_new")

DISCLAIMER = (
    "Выводы сформированы автоматически на основании сопоставления текстов представленных документов. "
    "Они носят рекомендательный характер и требуют проверки ответственным сотрудником. "
    "Каждый вывод сопровождается ссылкой на пункт документа и дословной цитатой; "
    "выводы без подтверждённой цитаты в заключение не включались."
)


def _src(c: dict) -> str:
    return f"{c['doc']}, п. {c['id']}"


def _quote(t: str, n: int = 220) -> str:
    return t if len(t) <= n else t[: n - 1].rstrip() + "…"


COI_LABELS = {"risk": "риск конфликта", "safeguard": "защитная норма", "mention": "упоминание"}


def coi_label(c: dict) -> str:
    v = c.get("verdict")
    return f" [{COI_LABELS[v['verdict']]}]" if v else ""


def _abolished(result: dict) -> set[str]:
    return set(result["structure"].get("positions_removed", []))


def _is_lead(c: dict) -> bool:
    if "lead" in c:
        return bool(c["lead"])
    return c.get("modality") == "none" and c["text"].rstrip().endswith(":")


def rep(f: dict, side: str = "a") -> dict | None:
    """Представитель связи: первый пункт-функция, а не заголовок блока. При связи 2:1
    старая сторона может начинаться с заголовка «Директор направления:», и без этого
    правила вывод получил бы владельца None и попал бы не в тот раздел."""
    cl = f.get(side) or []
    for c in cl:
        if not _is_lead(c):
            return c
    return cl[0] if cl else None


def is_empty(c: dict) -> bool:
    """Пункт без букв: в исходнике стоит одна точка с запятой. Это дефект документа,
    а не функция, и в потери он не идёт."""
    return not any(ch.isalpha() for ch in c["text"])


def data_issues(result: dict) -> list[dict]:
    out = []
    for f in result["findings"]:
        for side in ("a", "b"):
            for c in f.get(side) or []:
                if c["section"] in {"4", "5"} and c.get("id", "").count(".") >= 1 and is_empty(c):
                    out.append(c)
    seen, uniq = set(), []
    for c in out:
        if (c["doc"], c["id"]) not in seen:
            seen.add((c["doc"], c["id"]))
            uniq.append(c)
    return uniq


def losses(result: dict) -> list[dict]:
    """Потери функций у действующих владельцев. Функции упразднённой должности это
    перераспределение, они идут отдельным подразделом, см. redistributed()."""
    ab = _abolished(result)
    out = []
    for f in result["findings"]:
        a = rep(f)
        if f["kind"] in LOSS_KINDS and a and not _is_lead(a) and a["section"] in {"4", "5"} and (a["owner"] or "") not in ab and not is_empty(a):
            out.append(f)
    return out


def redistributed(result: dict) -> list[dict]:
    ab = _abolished(result)
    out = []
    for f in result["findings"]:
        a = rep(f)
        if f["kind"] in LOSS_KINDS + MOVE_KINDS and a and a["section"] in {"4", "5"} and (a["owner"] or "") in ab:
            out.append(f)
    return out


def function_table(result: dict) -> list[dict]:
    """Таблица сопоставления функций: только пункты разделов с функциями ролей."""
    rows = []
    for f in result["findings"]:
        if f["kind"] == "unchanged":
            continue
        a = rep(f) if f["a"] else None
        b = rep(f, "b")
        sec = (a or b or {}).get("section")
        if sec not in {"4", "5"}:
            continue
        rows.append(
            {
                "kind": f["title"],
                "before": f"{a['id']}: {_quote(a['text'], 110)}" if a else "—",
                "before_owner": (a or {}).get("owner") or "",
                "after": f"{b['id']}: {_quote(b['text'], 110)}" if b else "—",
                "after_owner": (b or {}).get("owner") or "",
                "sim": f["sim"],
            }
        )
    return rows


def lost_part(f: dict) -> str:
    d = f.get("llm") or {}
    if d.get("coverage") == "partial" and d.get("lost_part"):
        return d["lost_part"]
    return ""


# ---------------------------------------------------------------- рекомендации


def recommend(result: dict, model: str = "gpt-4.1-mini") -> dict[str, str]:
    """ref -> рекомендация. ref это идентификатор старого пункта для потерь и первый
    идентификатор кластера для дублей."""
    items = []
    for f in losses(result):
        a, b = rep(f), (rep(f, "b"))
        items.append(
            {
                "ref": a["id"],
                "kind": f["title"],
                "owner": a["owner"],
                "old": _quote(a["text"], 260),
                "new": _quote(b["text"], 260) if b else "",
                "lost": lost_part(f),
                "new_owner": (b or {}).get("owner") or "",
            }
        )
    for g in result["duplicates"]:
        v = g.get("verdict") or {}
        items.append(
            {
                "ref": g["clauses"][0]["id"],
                "kind": "Дублирование",
                "owner": "; ".join(sorted({c["owner"] or "" for c in g["clauses"]})),
                "old": "",
                "new": "; ".join(f"{c['id']}: {_quote(c['text'], 90)}" for c in g["clauses"][:4]),
                "lost": v.get("shared_function", ""),
                "new_owner": "",
            }
        )
    for c in result["coi"]:
        if c["id"].count(".") == 1:
            items.append(
                {
                    "ref": c["id"],
                    "kind": "Конфликт интересов",
                    "owner": c["owner"] or "",
                    "old": "",
                    "new": _quote(c["text"], 300),
                    "lost": "",
                    "new_owner": "",
                }
            )

    if not items:
        return {}
    if not os.environ.get("OPENAI_API_KEY"):
        return {it["ref"]: _template(it) for it in items}

    from .llm import make_client

    client = make_client()
    schema = {
        "type": "object",
        "additionalProperties": False,
        "properties": {
            "items": {
                "type": "array",
                "items": {
                    "type": "object",
                    "additionalProperties": False,
                    "properties": {"ref": {"type": "string"}, "recommendation": {"type": "string"}},
                    "required": ["ref", "recommendation"],
                },
            }
        },
        "required": ["items"],
    }
    user = (
        "Ниже выводы сравнения двух редакций положения о подразделении. Для каждого дай одну короткую "
        "рекомендацию ответственному сотруднику, одно-два предложения, что проверить или как устранить: "
        "вернуть функцию, явно закрепить за владельцем, убрать дублирование, зафиксировать меры по конфликту "
        "интересов. Не выдумывай пункты и факты, опирайся только на переданные тексты. "
        "Верни рекомендацию для каждого ref ровно один раз.\n\n" + json.dumps(items, ensure_ascii=False)
    )
    kwargs = {} if model.startswith(("gpt-5", "o")) else {"temperature": 0}
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": "Ты аналитик организационных документов. Пиши по-русски, кратко, по делу."},
            {"role": "user", "content": user},
        ],
        response_format={"type": "json_schema", "json_schema": {"name": "recs", "strict": True, "schema": schema}},
        **kwargs,
    )
    data = json.loads(resp.choices[0].message.content)
    valid = {it["ref"] for it in items}
    out = {d["ref"]: d["recommendation"] for d in data["items"] if d["ref"] in valid}
    for it in items:
        out.setdefault(it["ref"], _template(it))
    return out


def _template(it: dict) -> str:
    if it["kind"] == "Дублирование":
        return "Проверить, намеренно ли функция закреплена за несколькими ролями; при необходимости оставить одного владельца и указать порядок взаимодействия остальных."
    if it["kind"] == "Конфликт интересов":
        return "Проверить достаточность мер по исключению конфликта интересов и зафиксировать их исполнение в отчётности."
    if it["kind"] == "Снято с владельца":
        return f"Подтвердить, что функция намеренно снята с «{it['owner']}» и закреплена за «{it['new_owner']}»; иначе вернуть её в перечень функций владельца."
    if it["lost"]:
        return f"Проверить, намеренно ли из функции исключено «{it['lost']}»; при необходимости вернуть формулировку в новую редакцию."
    return f"Проверить, намеренно ли функция исключена у «{it['owner']}»; если нет, вернуть её в новую редакцию или явно закрепить за другим подразделением."


# ---------------------------------------------------------------- сборка


def build_conclusion(result: dict, with_recommendations: bool = True, model: str = "gpt-4.1-mini") -> dict:
    recs = recommend(result, model=model) if with_recommendations else {}
    return {"date": date.today().isoformat(), "result": result, "recs": recs, "losses": losses(result), "table": function_table(result)}


def to_markdown(conc: dict) -> str:
    r, recs = conc["result"], conc["recs"]
    m, s = r["meta"], r["structure"]
    w = []
    w.append("# Аналитическое заключение по сопоставлению редакций\n")
    w.append(f"Дата: {conc['date']}. Документы «до»: {', '.join(m['before'])}. Документы «после»: {', '.join(m['after'])}.\n")
    w.append(f"_{DISCLAIMER}_\n")

    w.append("## 1. Изменения организационной структуры\n")
    nb, na = s["names_before"], s["names_after"]
    for a in s["departments_created"]:
        w.append(f"- **Создано:** {na.get(a, a)} ({a}). Источник: после, п. {s['sources_after'].get(a, '3.4')}")
    for a in s["departments_removed"]:
        w.append(f"- **Упразднено:** {nb.get(a, a)} ({a}). Источник: до, п. {s['sources_before'].get(a, '3.4')}")
    for a in s["departments_kept"]:
        w.append(f"- **Сохранено:** {na.get(a, a)} ({a})")
    for p in s["positions_removed"]:
        w.append(f"- **Упразднена должность:** {p} (до, п. 3.5)")
    for p in s["positions_created"]:
        w.append(f"- **Введена должность:** {p} (после, п. 3.5)")
    rem = [x for x in s.get("unit_positions_removed", []) if x[0] != "Главный аудитор"]
    add = [x for x in s.get("unit_positions_created", []) if x[0] != "Главный аудитор"]
    if rem or add:
        w.append("- Изменения состава должностей внутри подразделений:")
        for boss, p in rem:
            w.append(f"  - исключена «{p}» у {boss} (до, п. {s['sources_before'].get(boss + ':' + p, '3')})")
        for boss, p in add:
            w.append(f"  - добавлена «{p}» у {boss} (после, п. {s['sources_after'].get(boss + ':' + p, '3')})")
    w.append("")

    redis = redistributed(r)
    if redis:
        w.append("### Перераспределение функций упразднённой должности\n")
        for f in redis:
            a = rep(f)
            b = rep(f, "b")
            w.append(
                f"- «{_quote(a['text'], 110)}» ({_src(a)}, {a['owner']}) → {(_src(b) + ', ' + (b['owner'] or '')) if b else 'получатель не найден'}. {f['title']}."
            )
        w.append("")

    w.append("## 2. Потери функций\n")
    if not conc["losses"]:
        w.append("Потерь функций не выявлено.\n")
    for f in conc["losses"]:
        a = rep(f)
        b = rep(f, "b")
        w.append(f"### {f['title']}: {a['owner'] or 'без владельца'}, п. {a['id']}\n")
        w.append(f"**Было** ({_src(a)}): {a['text']}\n")
        if b:
            w.append(f"**Стало** ({_src(b)}): {b['text']}\n")
        elif f["kind"] == "lost":
            w.append("**Стало:** в новой редакции пункта с этой функцией нет.\n")
            cand = f.get("candidate")
            if cand:
                w.append(
                    f"Ближайший по тексту пункт новой редакции, функцию не покрывает ({_src(cand)}, близость {f['sim']:.2f}): «{_quote(cand['text'], 160)}»\n"
                )
        lp = lost_part(f)
        if lp:
            w.append(f"**Пропало:** «{lp}»\n")
        d = f.get("llm") or {}
        if d.get("reason"):
            w.append(f"Обоснование: {d['reason']}\n")
        if a["id"] in recs:
            w.append(f"**Рекомендация:** {recs[a['id']]}\n")

    w.append("## 3. Переносы, слияния и смена владельца\n")
    ab = _abolished(r)
    moves = [f for f in r["findings"] if f["kind"] in MOVE_KINDS and f["a"] and rep(f)["section"] in {"4", "5"} and (rep(f)["owner"] or "") not in ab]
    if not moves:
        w.append("Не выявлено.\n")
    for f in moves:
        a = rep(f)
        b = rep(f, "b")
        w.append(
            f"- **{f['title']}**: {_src(a)} → {_src(b) if b else '—'}. {a['owner'] or ''} → {(b or {}).get('owner') or ''}. «{_quote(a['text'], 120)}»"
        )
    w.append("")

    w.append("## 4. Новые функции\n")
    news = [f for f in r["findings"] if f["kind"] in NEW_KINDS and f["b"] and f["b"][0]["section"] in {"4", "5"}]
    if not news:
        w.append("Не выявлено.\n")
    for f in news:
        b = f["b"][0]
        w.append(f"- {_src(b)}, {b['owner'] or 'без владельца'}: «{_quote(b['text'], 160)}»")
    w.append("")

    w.append("## 5. Дублирование функций\n")
    if not r["duplicates"]:
        w.append("Не выявлено.\n")
    for g in r["duplicates"]:
        v = g.get("verdict") or {}
        title = {"same": "Одна и та же функция", "overlap": "Пересечение функций"}.get(v.get("verdict"), "Похожие функции")
        w.append(f"### {title}: {v.get('shared_function') or _quote(g['clauses'][0]['text'], 80)}\n")
        for c in g["clauses"]:
            w.append(f"- {_src(c)}, {c['owner']}: «{_quote(c['text'], 160)}»")
        if v.get("reason"):
            w.append(f"\nОбоснование: {v['reason']}\n")
        if g["clauses"][0]["id"] in recs:
            w.append(f"**Рекомендация:** {recs[g['clauses'][0]['id']]}\n")

    w.append("## 6. Потенциальные конфликты интересов\n")
    if not r["coi"]:
        w.append("Не выявлено.\n")
    for c in r["coi"]:
        w.append(f"- {_src(c)}, {c['owner'] or 'без владельца'}{coi_label(c)}: «{_quote(c['text'], 240)}»")
        if c.get("verdict") and c["verdict"].get("quotes_ok") and c["verdict"]["verdict"] != "mention":
            w.append(f"  Основание: «{_quote(c['verdict']['quote'], 200)}». {c['verdict']['reason']}")
        if c["id"] in recs:
            w.append(f"  **Рекомендация:** {recs[c['id']]}")
    w.append("")

    w.append("## 7. Таблица сопоставления функций (изменившиеся пункты)\n")
    w.append("| Вид | До | Владелец до | После | Владелец после |")
    w.append("|---|---|---|---|---|")
    for row in conc["table"]:
        w.append(f"| {row['kind']} | {row['before']} | {row['before_owner']} | {row['after']} | {row['after_owner']} |")
    w.append("")

    issues = data_issues(r)
    if issues:
        w.append("## Замечания к исходным документам\n")
        for c in issues:
            w.append(f"- {_src(c)}: пункт пустой, в документе стоит «{c['text']}». Это дефект исходника, в анализе функций не участвует.")
        w.append("")

    w.append("## 8. Сведения о проверке\n")
    w.append(f"- Пунктов разобрано: до {m['clauses_before']}, после {m['clauses_after']}. Режим близости: {m['mode']}.")
    if m["verify"]:
        w.append(
            f"- Спорные места проверены моделью {m['model']}: {m['llm_calls']} обращений, отклонено вердиктов с неподтверждённой цитатой: {m['bad_quotes_blocked']}."
        )
    else:
        w.append("- Проверка моделью не выполнялась, выводы основаны только на близости текстов.")
    w.append(f"- Время анализа: {m['seconds']} с.\n")
    return "\n".join(w)


def to_docx(conc: dict, path: str) -> str:
    from docx import Document
    from docx.shared import Pt

    md = to_markdown(conc)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows
        if not table_rows:
            return
        t = doc.add_table(rows=0, cols=len(table_rows[0]))
        t.style = "Table Grid"
        for i, row in enumerate(table_rows):
            cells = t.add_row().cells
            for c, val in zip(cells, row, strict=False):
                c.text = val
                if i == 0:
                    for p in c.paragraphs:
                        for run in p.runs:
                            run.bold = True
        table_rows = []

    for line in md.split("\n"):
        if line.startswith("|"):
            if set(line.replace("|", "").strip()) <= {"-", " "}:
                continue
            table_rows.append([x.strip() for x in line.strip("|").split("|")])
            continue
        flush_table()
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- ") or line.startswith("  - "):
            p = doc.add_paragraph(style="List Bullet")
            _runs(p, line.lstrip(" -"))
        elif line.startswith("_") and line.endswith("_"):
            p = doc.add_paragraph()
            r = p.add_run(line.strip("_"))
            r.italic = True
        elif line.strip():
            p = doc.add_paragraph()
            _runs(p, line)
    flush_table()
    doc.save(path)
    return path


def _runs(p, text: str):
    parts = text.split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = i % 2 == 1

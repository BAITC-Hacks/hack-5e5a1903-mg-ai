"""Организационная структура как отдельная сущность.

Структура может прийти двумя путями: из раздела 3 положения (docx или pdf) или из
таблицы Excel «Подразделение / Аббревиатура / Должность / Подчиняется». Обе дороги
приводят к одному объекту Structure, и сравнение редакций делается по нему, а не по
пунктам. Это и есть тест парсера Excel: Structure из xlsx обязана совпасть со Structure
из docx для того же документа.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from .parse import Clause, departments, positions_under_chief

SUBORD = re.compile(r"^(?:Директору|Главному аудитору)\s*([А-ЯЁа-яё\s]*?)\s*подчиняются", re.I)


@dataclass
class Structure:
    units: dict[str, str] = field(default_factory=dict)  # аббревиатура -> полное имя
    positions: dict[str, list[str]] = field(default_factory=dict)  # начальник -> должности
    sources: dict[str, str] = field(default_factory=dict)  # что угодно -> пункт или строка-источник

    @property
    def chief_positions(self) -> set[str]:
        return set(self.positions.get("Главный аудитор", []))

    def all_positions(self) -> set[tuple[str, str]]:
        return {(boss, p) for boss, ps in self.positions.items() for p in ps}


def from_clauses(clauses: list[Clause]) -> Structure:
    s = Structure(units=departments(clauses))
    for abbr in s.units:
        s.sources[abbr] = next((c.id for c in clauses if c.num == "3.4" and abbr in c.text), "3.4")
    s.positions["Главный аудитор"] = [p for p in positions_under_chief(clauses)]
    by_id = {c.id: c for c in clauses}
    for c in clauses:
        if c.section != "3" or not c.num or c.num in {"3.4", "3.5"} or c.id == c.num:
            continue
        lead = by_id.get(c.num)
        if not lead or not lead.is_lead:
            continue
        m = SUBORD.match(lead.text)
        if not m or lead.text.startswith("Главному"):
            continue
        boss = f"Директор {m.group(1).strip()}".strip()
        s.positions.setdefault(boss, []).append(c.text.rstrip("."))
        s.sources[f"{boss}:{c.text.rstrip('.')}"] = c.id
    return s


def from_xlsx(path: str | Path) -> Structure:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return Structure()
    header = [str(h or "").strip().lower() for h in rows[0]]

    def col(*names):
        for n in names:
            for i, h in enumerate(header):
                if n in h:
                    return i
        return None

    c_unit, c_abbr, c_pos, c_boss, c_src = (
        col("подразделение", "департамент", "unit"),
        col("аббревиатура", "сокращ", "abbr"),
        col("должность", "position"),
        col("подчиня", "руковод", "boss"),
        col("источник", "source"),
    )
    s = Structure()
    for r_i, row in enumerate(rows[1:], start=2):

        def get(i, row=row):
            return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ""

        unit, abbr, pos, boss, src = get(c_unit), get(c_abbr), get(c_pos), get(c_boss), get(c_src)
        if unit and abbr and not pos:
            s.units[abbr] = unit
            s.sources[abbr] = src or f"{Path(path).name}:строка {r_i}"
        elif pos:
            boss = boss or "Главный аудитор"
            s.positions.setdefault(boss, []).append(pos)
            s.sources[f"{boss}:{pos}"] = src or f"{Path(path).name}:строка {r_i}"
    return s


def structure_changes(a: Structure, b: Structure) -> dict:
    pa, pb = a.chief_positions, b.chief_positions
    all_a, all_b = a.all_positions(), b.all_positions()
    return {
        "departments_created": sorted(set(b.units) - set(a.units)),
        "departments_removed": sorted(set(a.units) - set(b.units)),
        "departments_kept": sorted(set(a.units) & set(b.units)),
        "positions_created": sorted(pb - pa),
        "positions_removed": sorted(pa - pb),
        "positions_kept": sorted(pa & pb),
        "unit_positions_created": sorted(all_b - all_a),
        "unit_positions_removed": sorted(all_a - all_b),
        "names_before": a.units,
        "names_after": b.units,
        "sources_before": a.sources,
        "sources_after": b.sources,
    }
